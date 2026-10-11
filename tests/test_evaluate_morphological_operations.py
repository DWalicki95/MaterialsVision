"""Tests for scoring the classical predictions with the model evaluation.

The end-to-end tests write a miniature dataset - images, instance
masks, a manifest, a split and size classes - and classical
predictions in the directory layout Fiji left them in, then run the
real entry point over them. That covers the seam that matters: the
annotation the classical masks are scored against is the one the
sample source hands the models, cropped the same way.
"""
import json

import numpy as np
import pandas as pd
import pytest
import tifffile
from PIL import Image

import scripts.compare_classical_baseline as compare_cli
from materials_vision.experiments.evaluate_morphological_operations import \
    data_loading
from materials_vision.experiments.evaluate_morphological_operations import \
    main as cli
from scripts.summarize_cv import FoldScores, SummaryError

PredictionLoadError = data_loading.PredictionLoadError
as_instance_labels = data_loading.as_instance_labels
find_samples = data_loading.find_samples
image_id_from_sample_dir = data_loading.image_id_from_sample_dir

FRAME = (40, 32)                      # (height, width)
PANEL_ROWS = 10
EXPORT_SUFFIX = "_jpg.rf.0123abcd"

# image_id, formulation, material, microscope, split, panel
IMAGES = [
    ("AS1_0", "AS1", "AS", "M1", "train", False),
    ("AS1_1", "AS1", "AS", "M1", "train", False),
    ("K1_0", "K1", "K", "M2", "train", True),
    ("AS3_0", "AS3", "AS", "M1", "test", False),
]


def _gt_labels(panel: bool) -> np.ndarray:
    """Two pores in the content region; a third inside the panel."""
    labels = np.zeros(FRAME, dtype=np.uint16)
    labels[2:12, 2:12] = 1
    labels[14:26, 14:28] = 2
    if panel:
        labels[FRAME[0] - 8:FRAME[0] - 2, 2:12] = 3
    return labels


@pytest.fixture
def dataset(tmp_path):
    """Images, masks, manifest, split and size classes on disk."""
    images, masks = tmp_path / "images", tmp_path / "masks"
    images.mkdir()
    masks.mkdir()
    manifest_rows, split_rows = [], []
    for image_id, formulation, material, microscope, split, panel in IMAGES:
        Image.fromarray(np.full(FRAME, 100, dtype=np.uint8), mode="L").save(
            images / f"{image_id}.png"
        )
        labels = _gt_labels(panel)
        tifffile.imwrite(masks / f"{image_id}_masks.tif", labels)
        content_height = FRAME[0] - PANEL_ROWS if panel else FRAME[0]
        manifest_rows.append({
            "image_id": image_id,
            "load_crop_bbox": f"0,0,{FRAME[1]},{content_height}",
            "mask_path": str(masks / f"{image_id}_masks.tif"),
            "width_px": FRAME[1], "height_px": FRAME[0],
            "n_channels": 1, "channels_identical": None,
            "n_instances": int(labels.max()), "q_max_i": 1.0,
        })
        split_rows.append({
            "image_id": image_id, "formulation": formulation,
            "material": material, "microscope": microscope,
            "scale_bin": "coarse", "pixel_size_um": 3.0,
            "n_instances": int(labels.max()),
            "source_path": str(images / f"{image_id}.png"),
            "split": split, "used": True,
        })
    manifest_csv = tmp_path / "manifest.csv"
    pd.DataFrame(manifest_rows).to_csv(manifest_csv, index=False)
    split_csv = tmp_path / "split_test.csv"
    pd.DataFrame(split_rows).to_csv(split_csv, index=False)
    (tmp_path / "split_test_metadata.json").write_text(
        json.dumps({"split_id": "split_test"})
    )
    size_bins = tmp_path / "size_bins.json"
    size_bins.write_text(json.dumps({"bins": {
        "unit": "um2", "edges_um2": [500.0, 1000.0, 1500.0],
        "n_calibration_instances": 100,
    }}))
    return {
        "root": tmp_path / "predictions", "manifest": manifest_csv,
        "split": split_csv, "size_bins": size_bins,
        "out_dir": tmp_path / "out",
    }


def _write_prediction(root, image_id, variant, labels, group=None):
    """Store a prediction as Fiji did: float32, in the nested layout."""
    sample = f"{image_id}{EXPORT_SUFFIX}"
    group = group or image_id.split("_")[0]
    target = root / group / sample / "predicted_image" / variant
    target.mkdir(parents=True, exist_ok=True)
    tifffile.imwrite(target / f"{sample}_masks.tif",
                     labels.astype(np.float32))


def _predict(dataset, image_ids):
    """A perfect variant, and one that misses the second pore."""
    for image_id, *_, panel in [row for row in IMAGES
                                if row[0] in image_ids]:
        perfect = _gt_labels(panel)
        _write_prediction(dataset["root"], image_id, "perfect", perfect)
        partial = perfect.copy()
        partial[partial == 2] = 0
        _write_prediction(dataset["root"], image_id, "partial", partial)


def _run(dataset, *extra):
    args = cli.parse_args([
        "--predictions-root", str(dataset["root"]),
        "--manifest", str(dataset["manifest"]),
        "--split", str(dataset["split"]),
        "--size-bins", str(dataset["size_bins"]),
        "--out-dir", str(dataset["out_dir"]),
        "--variants", "perfect", "partial", "--no-figures", *extra,
    ])
    return cli.run(args)


def _scores(dataset):
    return json.loads(
        (dataset["out_dir"] / cli.SCORES_FILE).read_text()
    )


def test_image_id_is_recovered_from_the_export_name():
    assert image_id_from_sample_dir(
        "AS1_40_14_jpg.rf.675e59e3641e9b985aa3282dfc6f641f"
    ) == "AS1_40_14"
    assert image_id_from_sample_dir("VAB1 prostopadla_m001") == (
        "VAB1 prostopadla_m001"
    )


def test_whole_number_floats_become_dense_labels_without_mutation():
    stored = np.array([[0, 5], [9, 5]], dtype=np.float32)
    before = stored.copy()

    labels = as_instance_labels(stored)

    assert labels.dtype == np.int32
    np.testing.assert_array_equal(labels, [[0, 1], [2, 1]])
    np.testing.assert_array_equal(stored, before)


@pytest.mark.parametrize("stored", [
    np.array([[0.0, 1.5]]),
    np.array([[0.0, -1.0]]),
    np.array([[0.0, np.nan]]),
    np.zeros((2, 2, 3)),
])
def test_arrays_that_are_not_label_images_are_refused(stored):
    with pytest.raises(PredictionLoadError):
        as_instance_labels(stored)


def test_a_variant_missing_on_one_image_is_refused(dataset):
    _predict(dataset, {"AS1_0", "AS1_1"})
    missing = (dataset["root"] / "AS1" / f"AS1_1{EXPORT_SUFFIX}"
               / "predicted_image" / "partial")
    for path in missing.iterdir():
        path.unlink()

    with pytest.raises(PredictionLoadError, match="expected one mask"):
        find_samples(dataset["root"], ("perfect", "partial"))


def test_two_directories_for_one_image_are_refused(dataset):
    _predict(dataset, {"AS1_0"})
    labels = _gt_labels(False)
    for variant in ("perfect", "partial"):
        _write_prediction(dataset["root"], "AS1_0", variant, labels,
                          group="copy")

    with pytest.raises(PredictionLoadError, match="both hold"):
        find_samples(dataset["root"], ("perfect", "partial"))


def test_scores_come_from_the_annotation_the_models_were_scored_on(dataset):
    _predict(dataset, {"AS1_0", "AS1_1", "K1_0"})

    assert _run(dataset) == cli.EXIT_OK

    scores = _scores(dataset)
    perfect = scores["perfect"]["overall"]
    partial = scores["partial"]["overall"]
    # The pore inside K1_0's panel is cropped away on both sides, so
    # six annotated pores remain and the perfect variant matches all.
    assert perfect["n_gt"] == partial["n_gt"] == 6
    assert perfect["f1"] == pytest.approx(1.0)
    assert partial["recall"] == pytest.approx(0.5)
    assert partial["precision"] == pytest.approx(1.0)
    assert [row["label"] for row in scores["perfect"]["per_image"]] == [
        "AS1_0", "AS1_1", "K1_0",
    ]
    assert {row["label"] for row in scores["perfect"]["per_material"]} == {
        "material=AS", "material=K",
    }


def test_images_in_test_need_an_explicit_unlock(dataset):
    _predict(dataset, {"AS1_0", "AS3_0"})

    assert _run(dataset) == cli.EXIT_FAILED
    assert not (dataset["out_dir"] / cli.SCORES_FILE).exists()

    assert _run(dataset, "--unlock-test") == cli.EXIT_OK
    assert _scores(dataset)["perfect"]["overall"]["n_images"] == 2


def test_a_prediction_on_another_grid_is_refused_not_resized(dataset):
    _predict(dataset, {"AS1_0"})
    sample = dataset["root"] / "AS1" / f"AS1_0{EXPORT_SUFFIX}"
    small = np.zeros((FRAME[0] // 2, FRAME[1] // 2), dtype=np.float32)
    small[1:5, 1:5] = 1
    for variant in ("perfect", "partial"):
        tifffile.imwrite(
            sample / "predicted_image" / variant
            / f"AS1_0{EXPORT_SUFFIX}_masks.tif", small,
        )

    assert _run(dataset) == cli.EXIT_FAILED


def test_comparison_reads_the_report_the_evaluation_writes(dataset):
    _predict(dataset, {"AS1_0", "K1_0"})
    assert _run(dataset) == cli.EXIT_OK

    rows = compare_cli.classical_rows(_scores(dataset))

    by_method = {row["method"]: row for row in rows}
    perfect = by_method["classical perfect"]
    assert perfect["f1"] == pytest.approx(1.0)
    assert perfect["n_images"] == 2
    assert perfect["mean_pair_iou"] == pytest.approx(1.0)
    assert perfect["classical_only"]["boundary_f1"] == pytest.approx(1.0)
    assert set(perfect["per_image"]) == {"AS1_0", "K1_0"}
    assert by_method["classical partial"]["size_bin_recall"]["Q4"] == (
        pytest.approx(0.0)
    )
    report = compare_cli.format_markdown(rows)
    assert "| AS1_0 | 1.000 | 0.667 |" in report


def test_classical_variants_on_different_pores_are_refused():
    def variant(n_gt):
        return {"overall": {"n_images": 2, "n_gt": n_gt}, "per_image": []}

    with pytest.raises(SummaryError, match="different populations"):
        compare_cli.classical_rows({"a": variant(10), "b": variant(11)})


def _section(tp, fp, fn, iou, bins, n_images=5, porosity=1.0,
             wasserstein=10.0):
    return {
        "label": "overall", "n_images": n_images, "n_gt": tp + fn,
        "n_pred": tp + fp, "n_true_positives": tp,
        "n_false_positives": fp, "n_false_negatives": fn,
        "n_merges": 0, "n_splits": 0, "mean_pair_iou": iou,
        "size_bin_recall": [
            {"label": label, "n_gt": n_gt, "n_matched": n_matched}
            for label, (n_gt, n_matched) in zip(("Q1", "Q2", "Q3", "Q4"),
                                                bins)
        ],
        "mean_porosity_error_pp": porosity,
        "wasserstein_um": wasserstein,
        "median_diameter_drift_um": -1.0, "iqr_drift_um": 2.0,
        "median_diameter_error": 0.1, "median_elongation_error": 0.2,
        "median_angle_error_deg": 5.0,
        "orientation": {"resultant_difference": 0.05,
                        "grid_resultant_difference": -0.02},
    }


def test_pair_overlap_and_size_classes_pool_over_every_pair():
    folds = [
        _section(10, 0, 0, 0.9, [(4, 2), (2, 2), (2, 2), (2, 2)]),
        _section(30, 0, 0, 0.5, [(6, 3), (8, 8), (8, 8), (8, 8)]),
    ]

    row = compare_cli.method_row("model", "cv", folds)

    assert row["mean_pair_iou"] == pytest.approx((10 * 0.9 + 30 * 0.5) / 40)
    assert row["size_bin_recall"]["Q1"] == pytest.approx(5 / 10)
    assert row["n_gt"] == 40


def test_porosity_pools_over_images_and_the_rest_averages_folds():
    bins = [(1, 1)] * 4
    folds = [
        _section(10, 0, 0, 0.9, bins, n_images=10, porosity=2.0,
                 wasserstein=10.0),
        _section(10, 0, 0, 0.9, bins, n_images=30, porosity=6.0,
                 wasserstein=float("nan")),
        _section(10, 0, 0, 0.9, bins, n_images=20, porosity=4.0,
                 wasserstein=30.0),
    ]

    errors = compare_cli.method_row("model", "cv", folds)[
        "descriptor_errors"
    ]

    assert errors["mean_porosity_error_pp"] == pytest.approx(
        (10 * 2.0 + 30 * 6.0 + 20 * 4.0) / 60
    )
    assert errors["wasserstein_um"] == pytest.approx(20.0)
    assert errors["orientation_grid_resultant_difference"] == (
        pytest.approx(-0.02)
    )


def test_sheet_holds_the_figures_the_comparison_computes(dataset):
    from openpyxl import load_workbook

    import scripts.export_comparison_xlsx as export_cli
    from materials_vision.evaluation import load_size_bins

    _predict(dataset, {"AS1_0", "K1_0"})
    assert _run(dataset) == cli.EXIT_OK
    model = compare_cli.method_row(
        "PEFT-SAM (frozen)", "cv",
        [_section(90, 10, 10, 0.9, [(5, 4)] * 4, wasserstein=12.0),
         _section(80, 20, 20, 0.8, [(5, 3)] * 4, wasserstein=float("nan"))],
    )
    rows = [model, *compare_cli.classical_rows(_scores(dataset))]
    destination = dataset["out_dir"] / "sheet.xlsx"

    export_cli.write_workbook(
        rows, load_size_bins(dataset["size_bins"]), destination
    )

    sheet = load_workbook(destination)["Zestawienie"]
    table = [[cell.value for cell in line] for line in sheet.iter_rows()]
    header = table[0]
    assert header[:3] == ["Metoda", "Populacja", "F1"]
    assert "Wasserstein średnic [µm] *" in header
    assert "Recall Q4" in header
    by_name = {line[0]: dict(zip(header, line)) for line in table[1:]}
    assert set(by_name) == {
        "PEFT-SAM (dotrenowany micro-sam)",
        "Klasyczna: perfect", "Klasyczna: partial",
    }
    peft = by_name["PEFT-SAM (dotrenowany micro-sam)"]
    assert peft["F1"] == pytest.approx(model["f1"])
    assert peft["Wasserstein średnic [µm] *"] == pytest.approx(12.0)
    assert peft["Populacja"] == "CV 2 foldów: 10 obr., 200 porów"
    assert by_name["Klasyczna: perfect"]["F1"] == pytest.approx(1.0)
    assert by_name["Klasyczna: partial"]["Recall Q4"] == pytest.approx(0.0)
    # Q1 holds no annotated pore in the fixture: left empty, not zero.
    assert by_name["Klasyczna: perfect"]["Recall Q1"] is None


def _figure_rows(dataset):
    """A model row and the two classical rows from the real report."""
    model = compare_cli.method_row(
        "PEFT-SAM (frozen)", "cv",
        [_section(90, 10, 10, 0.9, [(5, 4)] * 4),
         _section(80, 20, 20, 0.8, [(5, 3)] * 4)],
    )
    classical = compare_cli.classical_rows(_scores(dataset))
    names = {"classical perfect": "classical interactive_watershed",
             "classical partial": "classical marker_watershed"}
    for row in classical:
        row["method"] = names[row["method"]]
    return [model, *classical]


def test_every_figure_is_drawn_from_the_comparison_rows(dataset):
    import scripts.plot_classical_comparison as plot_cli
    from materials_vision.evaluation import load_size_bins

    _predict(dataset, {"AS1_0", "K1_0"})
    assert _run(dataset) == cli.EXIT_OK
    rows = _figure_rows(dataset)
    out_dir = dataset["out_dir"] / "figures"

    written = plot_cli.draw_all(
        rows, load_size_bins(dataset["size_bins"]), out_dir
    )

    assert len(written) == 12
    assert all(path.stat().st_size > 0 for path in written)
    assert rows[0]["f1_per_section"] == pytest.approx([0.9, 0.8])


def test_a_method_without_a_style_is_refused():
    import scripts.plot_classical_comparison as plot_cli

    with pytest.raises(ValueError, match="No figure style"):
        plot_cli.method_key({"method": "classical otsu"})


def test_numbers_are_written_the_polish_way():
    import scripts.plot_classical_comparison as plot_cli

    assert plot_cli.format_number(-0.0367, 3) == "−0,037"
    assert plot_cli.format_number(4.1, 1, signed=True) == "+4,1"
    assert plot_cli.format_number(float("nan"), 1) == "b.d."


def test_a_fold_without_a_descriptor_is_refused():
    fold = _section(10, 0, 0, 0.9, [(1, 1)] * 4)
    del fold["orientation"]

    with pytest.raises(SummaryError, match="orientation"):
        compare_cli.method_row("model", "cv", [fold])


def test_reference_refuses_an_incomplete_cross_validation(
    monkeypatch, tmp_path
):
    recipe = tmp_path / "recipe.json"
    recipe.write_text(json.dumps({"postprocessing": {
        "flow_threshold": 0.4, "cellprob_threshold": 0.0,
        "min_size_px": 15, "niter": 2000,
    }}))
    fold = FoldScores(fold=1, overall={}, per_material=[],
                      per_formulation=None, held_out={})
    monkeypatch.setattr(compare_cli, "load_cellpose_folds",
                        lambda *args: ([fold], [2, 3, 4, 5]))
    monkeypatch.setattr(compare_cli, "load_peft_sam_folds",
                        lambda *args: ([fold], []))
    args = compare_cli.parse_args(["--recipe", str(recipe)])

    with pytest.raises(SummaryError, match="Folds missing"):
        compare_cli.model_rows(args)
