#!/usr/bin/env python3
"""
Write the classical-baseline comparison as an Excel sheet for slides.

The figures are the ones ``compare_classical_baseline.py`` computes -
the models pooled over the five folds of ``cv5_v1``, the classical
variants over the images they were run on - read afresh on every run,
so the sheet cannot drift from the reports. Only a selection is
written, chosen for a presentation: instance F1, the pore count error,
recall per size class, and the errors in the descriptors a material
report is built from - porosity, the diameter distribution, per-pair
diameter and shape, and orientation.

One method per row. Columns whose model figure is a mean of the five
fold values rather than a pooled one carry a ``*`` in their header,
and a second sheet spells out what every column means, so a table
pasted on its own still says how to read it.

Examples
--------
    $ PYTHONPATH=. python scripts/export_comparison_xlsx.py
"""
import argparse
import logging
import math
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Optional, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

import scripts.compare_classical_baseline as comparison
from materials_vision.evaluation import (SIZE_BIN_LABELS, SizeBins,
                                         SizeBinsLoadError, load_size_bins)
from materials_vision.logging_config import setup_logging
from scripts.summarize_cv import SummaryError

logger = logging.getLogger(__name__)

EXIT_OK = 0
EXIT_FAILED = 1

DEFAULT_OUT = Path("outputs/morphological_evaluation/zestawienie_metod.xlsx")

DEFAULT_SIZE_BINS = Path("/home/dwalicki/dane/splits/size_bins_v1.json")

FOLD_MEAN_MARK = " *"

# Excel number formats. Signed figures show their sign, so that an
# over- and an under-estimate read differently at a glance.
SCORE = "0.000"
SIGNED_PERCENT = "+0.0%;-0.0%;0.0%"
PERCENT = "0.0%"
SIGNED_ONE_DECIMAL = "+0.0;-0.0;0.0"
ONE_DECIMAL = "0.0"
SIGNED_SCORE = "+0.000;-0.000;0.000"


@dataclass(frozen=True)
class Column:
    """One column of the sheet.

    Parameters
    ----------
    header : str
    path : tuple of str
        Keys leading to the figure in a comparison row.
    number_format : str
    fold_mean : bool
        Whether the models' figure is a mean of fold values.
    """

    header: str
    path: tuple[str, ...]
    number_format: str
    fold_mean: bool = False


COLUMNS = (
    Column("F1", ("f1",), SCORE),
    Column("Błąd liczby porów", ("pore_count_error",), SIGNED_PERCENT),
    *(
        Column(f"Recall {label}", ("size_bin_recall", label), SCORE)
        for label in SIZE_BIN_LABELS
    ),
    Column("Błąd porowatości [pp]",
           ("descriptor_errors", "mean_porosity_error_pp"),
           SIGNED_ONE_DECIMAL),
    Column("Wasserstein średnic [µm]",
           ("descriptor_errors", "wasserstein_um"), ONE_DECIMAL, True),
    Column("Przesunięcie mediany średnic [µm]",
           ("descriptor_errors", "median_diameter_drift_um"),
           SIGNED_ONE_DECIMAL, True),
    Column("Przesunięcie IQR średnic [µm]",
           ("descriptor_errors", "iqr_drift_um"), SIGNED_ONE_DECIMAL, True),
    Column("Błąd średnicy równoważnej",
           ("descriptor_errors", "median_diameter_error"), PERCENT, True),
    Column("Błąd wydłużenia",
           ("descriptor_errors", "median_elongation_error"), PERCENT, True),
    Column("Błąd kąta [°]",
           ("descriptor_errors", "median_angle_error_deg"), ONE_DECIMAL,
           True),
    Column("Δ koncentracji orientacji R",
           ("descriptor_errors", "orientation_resultant_difference"),
           SIGNED_SCORE, True),
    Column("Δ koncentracji siatki R4",
           ("descriptor_errors", "orientation_grid_resultant_difference"),
           SIGNED_SCORE, True),
)

LEADING_HEADERS = ("Metoda", "Populacja")


def display_name(row: dict[str, Any]) -> str:
    """The method's name as it should appear on a slide.

    Parameters
    ----------
    row : dict
        A row of ``compare_classical_baseline.compare``.

    Returns
    -------
    str
        A Polish label; the row's own name when it is not recognised.
    """
    method = row["method"]
    if method.startswith("PEFT-SAM"):
        return "PEFT-SAM (dotrenowany micro-sam)"
    if method.startswith("Cellpose-SAM"):
        return "Cellpose-SAM (dotrenowany)"
    if method.startswith("classical "):
        variant = method.removeprefix("classical ").replace("_", " ")
        return f"Klasyczna: {variant}"
    return method


def population(row: dict[str, Any]) -> str:
    """Which images the row rests on, in a few words.

    Parameters
    ----------
    row : dict

    Returns
    -------
    str
    """
    n_gt = f"{row['n_gt']:,}".replace(",", " ")
    counts = f"{row['n_images']} obr., {n_gt} porów"
    if "classical_only" in row:
        return f"{counts} (szacunek)"
    return f"CV {row['n_sections']} foldów: {counts}"


def cell_value(row: dict[str, Any], path: Sequence[str]) -> Optional[float]:
    """A figure for a cell; empty where it is undefined.

    Parameters
    ----------
    row : dict
    path : sequence of str

    Returns
    -------
    float or None
    """
    value: Any = row
    for key in path:
        value = value[key]
    if value is None or not math.isfinite(value):
        return None
    return float(value)


def notes(size_bins: SizeBins, rows: Sequence[dict[str, Any]]) -> list[str]:
    """What every column means, for the second sheet.

    Parameters
    ----------
    size_bins : SizeBins
    rows : sequence of dict

    Returns
    -------
    list of str
    """
    low, middle, high = (f"{edge:,.0f}".replace(",", " ")
                         for edge in size_bins.edges_um2)
    n_folds = max(
        (row["n_sections"] for row in rows if "classical_only" not in row),
        default=0,
    )
    return [
        "Populacje",
        "Modele: wyniki łączne walidacji krzyżowej cv5_v1; każdy obraz "
        "oceniony przez model, który nie widział jego formulacji. "
        "PEFT-SAM z postprocessingiem frozen, Cellpose-SAM z progami "
        "z przepisu (niter 2000).",
        "Metody klasyczne: szacunek na 10 obrazach AS (mikroskop M1), "
        "parametry dobierane ręcznie dla każdego obrazu. Maski dzielą "
        "cały kadr, więc ściany są przypisane do porów.",
        f"* Dla modeli średnia z {n_folds} foldów (mediany i odległości "
        f"rozkładów nie dają się połączyć z raportów foldów); dla metod "
        f"klasycznych wartość z jednego przebiegu.",
        "",
        "Definicje (znak: predykcja - adnotacja)",
        "F1: dopasowanie porów przy IoU masek >= 0.5, liczniki sumowane "
        "po obrazach.",
        "Błąd liczby porów: (N_pred - N_GT) / N_GT.",
        f"Recall Q1-Q4: odsetek porów adnotacji z danej klasy pola "
        f"dopasowanych przy IoU >= 0.5. Granice klas: {low}, {middle}, "
        f"{high} µm².",
        "Błąd porowatości: różnica udziału powierzchni porów w kadrze, "
        "w punktach procentowych, średnia po obrazach.",
        "Wasserstein, przesunięcie mediany i IQR: porównanie rozkładów "
        "średnicy równoważnej wszystkich porów (bez porów na brzegu "
        "kadru), w µm.",
        "Błąd średnicy równoważnej i wydłużenia: mediana |pred - GT| / "
        "GT po dopasowanych parach (bez porów na brzegu kadru).",
        "Błąd kąta: mediana różnicy orientacji osi par, w stopniach; "
        "pory o wydłużeniu < 1.2 pominięte.",
        "Δ R: różnica koncentracji orientacji wokół jednej osi; Δ R4: "
        "wokół pary osi prostopadłych (siatka pikseli). Czytać razem.",
    ]


def write_workbook(
    rows: Sequence[dict[str, Any]], size_bins: SizeBins, destination: Path
) -> Path:
    """Write the comparison sheet and its notes.

    Parameters
    ----------
    rows : sequence of dict
        Rows of ``compare_classical_baseline.compare``.
    size_bins : SizeBins
        The frozen size classes, for the notes.
    destination : Path

    Returns
    -------
    Path
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Zestawienie"
    headers = [
        *LEADING_HEADERS,
        *(column.header + (FOLD_MEAN_MARK if column.fold_mean else "")
          for column in COLUMNS),
    ]
    sheet.append(headers)
    for row in rows:
        sheet.append([
            display_name(row), population(row),
            *(cell_value(row, column.path) for column in COLUMNS),
        ])
    _format_table(sheet, len(rows))

    notes_sheet = workbook.create_sheet("Uwagi")
    for line in notes(size_bins, rows):
        notes_sheet.append([line])
    notes_sheet.column_dimensions["A"].width = 120

    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    logger.info("Wrote %s.", destination)
    return destination


def _format_table(sheet, n_rows: int) -> None:
    """Bold wrapped headers, number formats, widths, frozen first column."""
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center",
                                   horizontal="center")
    offset = len(LEADING_HEADERS) + 1
    for index, column in enumerate(COLUMNS, start=offset):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = 14
        for cell in sheet[letter][1:n_rows + 1]:
            cell.number_format = column.number_format
    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 30
    sheet.row_dimensions[1].height = 45
    sheet.freeze_panes = "B2"


def parse_args(argv: Optional[list[str]] = None) -> argparse.Namespace:
    """Read the command line.

    Parameters
    ----------
    argv : list of str, optional

    Returns
    -------
    argparse.Namespace
    """
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT)
    parser.add_argument("--size-bins", type=Path, default=DEFAULT_SIZE_BINS)
    parser.add_argument(
        "--classical-scores", type=Path,
        default=comparison.DEFAULT_CLASSICAL_SCORES,
    )
    parser.add_argument(
        "--recipe", type=Path, default=comparison.DEFAULT_RECIPE
    )
    parser.add_argument(
        "--cellpose-eval-dir", type=Path,
        default=comparison.DEFAULT_SAVE_ROOT / "eval",
    )
    parser.add_argument(
        "--peft-sam-eval-dir", type=Path,
        default=comparison.DEFAULT_EVAL_DIR,
    )
    parser.add_argument(
        "--folds-dir", type=Path, default=comparison.DEFAULT_FOLDS_DIR
    )
    return parser.parse_args(argv)


def main(argv: Optional[list[str]] = None) -> int:
    """Write the sheet.

    Returns
    -------
    int
        Process exit code.
    """
    args = parse_args(argv)
    setup_logging()
    try:
        rows = comparison.compare(args)["rows"]
        size_bins = load_size_bins(args.size_bins)
    except (SummaryError, SizeBinsLoadError, FileNotFoundError,
            KeyError) as error:
        logger.error("%s", error)
        return EXIT_FAILED
    write_workbook(rows, size_bins, args.out)
    return EXIT_OK


if __name__ == "__main__":
    sys.exit(main())
